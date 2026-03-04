"""
Training routes for Vanna semantic layer.
Handles business context and SQL training data uploads.
"""
import os
import shutil
import asyncio
from pathlib import Path
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from typing import Optional
from sqlalchemy import text
from database import AsyncSessionLocal
import logging

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/training", tags=["training"])

# Upload directory paths
UPLOAD_DIR = "/app/frontend/public/uploads"
SQL_UPLOAD_DIR = "/app/frontend/public/uploads/sql"


def extract_file_content(file_path: str) -> str:
    """Extract text content from uploaded file."""
    try:
        # For text-based files, read directly
        if file_path.endswith(('.txt', '.md', '.sql', '.csv')):
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        
        # For PDF files
        elif file_path.endswith('.pdf'):
            try:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    text_content = []
                    for page in reader.pages:
                        text_content.append(page.extract_text() or '')
                    return '\n'.join(text_content)
            except ImportError:
                # Fallback: read as binary and decode what we can
                with open(file_path, 'rb') as f:
                    content = f.read()
                    try:
                        return content.decode('utf-8', errors='ignore')
                    except:
                        return "[PDF content - install PyPDF2 for text extraction]"
        
        # For Word documents
        elif file_path.endswith(('.doc', '.docx')):
            try:
                import docx
                doc = docx.Document(file_path)
                return '\n'.join([para.text for para in doc.paragraphs])
            except ImportError:
                return "[Word document - install python-docx for text extraction]"
        
        # Default: try to read as text
        else:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
                
    except Exception as e:
        logger.error(f"Error extracting content from {file_path}: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to extract file content: {str(e)}")


@router.post("/business-context")
async def submit_business_context(
    file: Optional[UploadFile] = File(None),
    content: Optional[str] = Form(None)
):
    """
    Submit business context for Vanna training.
    Accepts either a file upload or text content (or both).
    """
    if not file and not content:
        raise HTTPException(status_code=400, detail="Either file or content must be provided")
    
    file_name = None
    file_path = None
    final_content = content or ""
    
    # Handle file upload
    if file and file.filename:
        try:
            # Ensure upload directory exists
            os.makedirs(UPLOAD_DIR, exist_ok=True)
            
            # Generate unique filename
            import uuid
            file_extension = os.path.splitext(file.filename)[1]
            unique_filename = f"{uuid.uuid4().hex}{file_extension}"
            full_path = os.path.join(UPLOAD_DIR, unique_filename)
            
            # Save the file
            with open(full_path, "wb") as buffer:
                shutil.copyfileobj(file.file, buffer)
            
            file_name = file.filename
            file_path = f"/uploads/{unique_filename}"  # Relative path for frontend access
            
            # Extract content from file
            extracted_content = extract_file_content(full_path)
            
            # Combine extracted content with any text content provided
            if content:
                final_content = f"{extracted_content}\n\n---\n\n{content}"
            else:
                final_content = extracted_content
                
            logger.info(f"File uploaded: {file_name} -> {file_path}")
            
        except Exception as e:
            logger.error(f"File upload error: {e}")
            raise HTTPException(status_code=500, detail=f"File upload failed: {str(e)}")
    
    # Save to database (is_sync=0 for new uploads)
    try:
        async with AsyncSessionLocal() as session:
            query = text("""
                INSERT INTO business_context (file_name, file_path, content, is_sync) 
                VALUES (:file_name, :file_path, :content, 0)
            """)
            await session.execute(query, {
                "file_name": file_name,
                "file_path": file_path,
                "content": final_content
            })
            await session.commit()
            
            # Get the inserted ID
            result = await session.execute(text("SELECT LAST_INSERT_ID()"))
            inserted_id = result.scalar()
            
        logger.info(f"Business context saved with ID: {inserted_id}")
        
        return {
            "success": True,
            "message": "Business context submitted successfully",
            "data": {
                "id": inserted_id,
                "file_name": file_name,
                "file_path": file_path,
                "content_length": len(final_content)
            }
        }
        
    except Exception as e:
        logger.error(f"Database error: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to save business context: {str(e)}")


@router.get("/business-context")
async def get_business_contexts():
    """Get all business context entries with is_sync status."""
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, file_name, file_path, 
                       LEFT(content, 600) as content_preview,
                       LENGTH(content) as content_length,
                       is_sync, created_at, updated_at 
                FROM business_context 
                ORDER BY created_at DESC
            """))
            rows = result.fetchall()
            
            contexts = []
            for row in rows:
                contexts.append({
                    "id": row[0],
                    "file_name": row[1],
                    "file_path": row[2],
                    "content_preview": row[3],
                    "content_length": row[4],
                    "is_sync": int(row[5]) if row[5] is not None else 0,
                    "created_at": str(row[6]) if row[6] else None,
                    "updated_at": str(row[7]) if row[7] else None
                })
            
            return {
                "success": True,
                "data": contexts,
                "total": len(contexts)
            }
            
    except Exception as e:
        logger.error(f"Error fetching business contexts: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to fetch business contexts: {str(e)}")


@router.post("/sql-training")
async def submit_sql_training(
    file: Optional[UploadFile] = File(None),
    content: Optional[str] = Form(None)
):
    """
    Submit SQL training data for Vanna training.
    Accepts either a file upload or text content (or both).
    """
    if not file and not content:
        raise HTTPException(status_code=400, detail="Either file or content must be provided")
    
    file_name = None
    file_path = None
    final_content = content or ""
    
    # Handle file upload
    if file and file.filename:
        try:
            # Ensure upload directory exists
            os.makedirs(SQL_UPLOAD_DIR, exist_ok=True)
            
            # Generate unique filename
            import uuid
            file_extension = os.path.splitext(file.filename)[1]
            unique_filename = f"{uuid.uuid4().hex}{file_extension}"
            full_path = os.path.join(SQL_UPLOAD_DIR, unique_filename)
            
            # Save the file
            with open(full_path, "wb") as buffer:
                shutil.copyfileobj(file.file, buffer)
            
            file_name = file.filename
            file_path = f"/uploads/sql/{unique_filename}"  # Relative path for frontend access
            
            # Extract content from file
            extracted_content = extract_file_content(full_path)
            
            # Combine extracted content with any text content provided
            if content:
                final_content = f"{extracted_content}\n\n---\n\n{content}"
            else:
                final_content = extracted_content
                
            logger.info(f"SQL training file uploaded: {file_name} -> {file_path}")
            
        except Exception as e:
            logger.error(f"SQL training file upload error: {e}")
            raise HTTPException(status_code=500, detail=f"File upload failed: {str(e)}")
    
    # Save to database
    try:
        async with AsyncSessionLocal() as session:
            query = text("""
                INSERT INTO sql_training (file_name, file_path, content, is_sync) 
                VALUES (:file_name, :file_path, :content, 0)
            """)
            await session.execute(query, {
                "file_name": file_name,
                "file_path": file_path,
                "content": final_content
            })
            await session.commit()
            
            # Get the inserted ID
            result = await session.execute(text("SELECT LAST_INSERT_ID()"))
            inserted_id = result.scalar()
            
        logger.info(f"SQL training data saved with ID: {inserted_id}")
        
        return {
            "success": True,
            "message": "SQL training data submitted successfully",
            "data": {
                "id": inserted_id,
                "file_name": file_name,
                "file_path": file_path,
                "content_length": len(final_content)
            }
        }
        
    except Exception as e:
        logger.error(f"Database error: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to save SQL training data: {str(e)}")


@router.get("/sql-training")
async def get_sql_training():
    """Get all SQL training entries with is_sync status."""
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, file_name, file_path, 
                       LEFT(content, 600) as content_preview,
                       LENGTH(content) as content_length,
                       is_sync, created_at, updated_at 
                FROM sql_training 
                ORDER BY created_at DESC
            """))
            rows = result.fetchall()
            
            entries = []
            for row in rows:
                entries.append({
                    "id": row[0],
                    "file_name": row[1],
                    "file_path": row[2],
                    "content_preview": row[3],
                    "content_length": row[4],
                    "is_sync": int(row[5]) if row[5] is not None else 0,
                    "created_at": str(row[6]) if row[6] else None,
                    "updated_at": str(row[7]) if row[7] else None
                })
            
            return {
                "success": True,
                "data": entries,
                "total": len(entries)
            }
            
    except Exception as e:
        logger.error(f"Error fetching SQL training data: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to fetch SQL training data: {str(e)}")


def parse_sql_training_content(content: str) -> list:
    """
    Parse SQL training content to extract question-SQL pairs.
    Supports formats:
    - JSON: [{"question": "...", "sql": "..."}, ...]
    - Q:/SQL: or Question:/SQL: or Query: (multiline)
    - Plain SQL: one or more statements separated by semicolons; each becomes a pair
      with a synthetic question derived from the SQL (first line or first 80 chars).
    """
    import re
    import json
    
    pairs = []
    content = (content or "").strip()
    if not content:
        return pairs
    
    # Try JSON format first
    try:
        data = json.loads(content)
        if isinstance(data, list):
            for item in data:
                if isinstance(item, dict) and 'question' in item and 'sql' in item:
                    q, s = item['question'], item['sql']
                    if q and s:
                        pairs.append({"question": str(q).strip(), "sql": _normalize_sql(str(s))})
            if pairs:
                return pairs
        if isinstance(data, dict) and 'question' in data and 'sql' in data:
            q, s = data['question'], data['sql']
            if q and s:
                return [{"question": str(q).strip(), "sql": _normalize_sql(str(s))}]
    except json.JSONDecodeError:
        pass
    
    # Try Q:/Question:/Query: ... SQL:/Query: ... format (multiline, flexible)
    pattern = r'(?:Q(?:uestion)?|Question)\s*:\s*(.+?)\s*(?:SQL|Query)\s*:\s*(.+?)(?=(?:Q(?:uestion)?|Question)\s*:|$)'
    matches = re.findall(pattern, content, re.DOTALL | re.IGNORECASE)
    for question, sql in matches:
        question = question.strip()
        sql = _normalize_sql(sql)
        if question and sql:
            pairs.append({"question": question, "sql": sql})
    if pairs:
        return pairs
    
    # Fallback: plain SQL — split by semicolons and treat each statement as a pair
    # (question = short description from first line / first 80 chars)
    statements = re.split(r';\s*', content)
    for stmt in statements:
        stmt = stmt.strip()
        if not stmt or len(stmt) < 10:
            continue
        # Skip comment-only or non-SQL lines
        lines = [l.strip() for l in stmt.splitlines() if l.strip() and not l.strip().startswith('--')]
        if not lines:
            continue
        first_line = lines[0].upper()
        if not (
            first_line.startswith('SELECT') or first_line.startswith('WITH') or
            first_line.startswith('INSERT') or first_line.startswith('UPDATE') or
            first_line.startswith('DELETE') or 'FROM ' in first_line or ' JOIN ' in first_line
        ):
            continue
        sql = _normalize_sql(stmt)
        question = _sql_to_question(stmt)
        pairs.append({"question": question, "sql": sql})
    
    return pairs


def _normalize_sql(sql: str) -> str:
    """Ensure SQL ends with exactly one semicolon."""
    if not sql or not sql.strip():
        return sql
    return sql.strip().rstrip(';').strip() + ';'


def _sql_to_question(sql: str, max_len: int = 80) -> str:
    """Derive a short question label from a SQL statement for plain-SQL training."""
    if not sql or not sql.strip():
        return "Execute query"
    # Use first non-comment line, truncated
    for line in sql.splitlines():
        line = line.strip()
        if line and not line.startswith('--'):
            if len(line) <= max_len:
                return line
            return line[: max_len - 3].rstrip() + "..."
    return "Execute query"


def _content_to_chunks(content: str) -> list:
    """Split content into chunks for Vanna training."""
    if not content or not content.strip():
        return []
    raw_chunks = content.split('\n\n')
    chunks = []
    for chunk in raw_chunks:
        chunk = chunk.strip()
        if chunk and len(chunk) > 20:
            chunks.append(chunk)
    if not chunks:
        chunks = [content.strip()]
    return chunks


async def _train_content_with_vanna(content: str) -> tuple:
    """
    Train Vanna with documentation content. Returns (trained_count, errors).
    """
    from services.semantic_layer import get_vanna_client
    vn = await get_vanna_client()
    if vn is None:
        return 0, ["Vanna client not initialized"]
    chunks = _content_to_chunks(content)
    trained_count = 0
    errors = []
    for i, chunk in enumerate(chunks):
        try:
            vn.train(documentation=chunk)
            trained_count += 1
            logger.info(f"Trained business context chunk {i+1}/{len(chunks)}")
        except Exception as e:
            error_msg = f"Failed to train chunk {i+1}: {str(e)}"
            logger.error(error_msg)
            errors.append(error_msg)
    return trained_count, errors


@router.post("/sync/business-context")
async def sync_business_context():
    """
    Sync all unsynced business context with Vanna training.
    Trains Vanna with all rows where is_sync=0, then sets is_sync=1 for those rows.
    """
    from config import VANNA_ENABLED
    
    if not VANNA_ENABLED:
        raise HTTPException(status_code=400, detail="Vanna is not enabled")
    
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, content, file_name 
                FROM business_context 
                WHERE is_sync = 0
                ORDER BY id ASC
            """))
            rows = result.fetchall()
        
        if not rows:
            return {
                "success": True,
                "message": "No unsynced business context to sync",
                "trained_count": 0,
                "synced_ids": []
            }
        
        from services.semantic_layer import get_vanna_client
        vn = await get_vanna_client()
        if vn is None:
            raise HTTPException(status_code=500, detail="Failed to get Vanna client")
        
        total_trained = 0
        all_errors = []
        synced_ids = []
        
        for row in rows:
            doc_id, content, file_name = row[0], row[1], row[2]
            if not content or not content.strip():
                continue
            trained_count, errors = await _train_content_with_vanna(content)
            total_trained += trained_count
            all_errors.extend(errors)
            synced_ids.append(doc_id)
        
        # Mark all synced rows as is_sync=1
        if synced_ids:
            placeholders = ",".join([str(i) for i in synced_ids])
            async with AsyncSessionLocal() as session:
                await session.execute(text(f"""
                    UPDATE business_context SET is_sync = 1 WHERE id IN ({placeholders})
                """))
                await session.commit()
        
        status = vn.get_training_status()
        return {
            "success": total_trained > 0,
            "message": f"Synced {len(synced_ids)} document(s) with Vanna",
            "trained_count": total_trained,
            "synced_ids": synced_ids,
            "errors": all_errors if all_errors else None,
            "training_status": status
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error syncing business context: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to sync business context: {str(e)}")


@router.post("/sync/business-context/{doc_id:int}")
async def sync_business_context_single(doc_id: int):
    """
    Sync a single business context document with Vanna training.
    Trains Vanna with that row's content and sets is_sync=1 for that row.
    """
    from config import VANNA_ENABLED
    
    if not VANNA_ENABLED:
        raise HTTPException(status_code=400, detail="Vanna is not enabled")
    
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, content, file_name FROM business_context WHERE id = :doc_id
            """), {"doc_id": doc_id})
            row = result.fetchone()
        
        if not row:
            raise HTTPException(status_code=404, detail="Business context document not found")
        
        content, file_name = row[1], row[2]
        if not content or not content.strip():
            return {
                "success": False,
                "message": "Document has no content to sync",
                "trained_count": 0
            }
        
        trained_count, errors = await _train_content_with_vanna(content)
        
        async with AsyncSessionLocal() as session:
            await session.execute(text("UPDATE business_context SET is_sync = 1 WHERE id = :doc_id"), {"doc_id": doc_id})
            await session.commit()
        
        from services.semantic_layer import get_vanna_client
        vn = await get_vanna_client()
        status = vn.get_training_status() if vn else {}
        
        return {
            "success": trained_count > 0,
            "message": f"Synced document (ID: {doc_id}, File: {file_name or 'text input'})",
            "trained_count": trained_count,
            "errors": errors if errors else None,
            "training_status": status
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error syncing business context doc {doc_id}: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to sync document: {str(e)}")


def _ensure_sql_faiss_files_exist():
    """
    Ensure sql_metadata.json exists so Vanna can load/append when files were deleted.
    When all SQL docs were deleted we remove sql_index.faiss and sql_metadata.json;
    the next sync must be able to create new training - Vanna may need metadata file to exist.
    """
    from config import VANNA_FAISS_PATH
    faiss_path = Path(VANNA_FAISS_PATH)
    faiss_path.mkdir(parents=True, exist_ok=True)
    sql_meta = faiss_path / "sql_metadata.json"
    if not sql_meta.exists():
        try:
            sql_meta.write_text("[]", encoding="utf-8")
            logger.info("Created empty sql_metadata.json for SQL training")
        except Exception as e:
            logger.warning(f"Could not create sql_metadata.json: {e}")


async def _train_sql_content_with_vanna(content: str) -> tuple:
    """
    Parse content and train Vanna with question-SQL pairs. Returns (trained_count, errors).
    """
    from services.semantic_layer import get_vanna_client
    vn = await get_vanna_client()
    if vn is None:
        return 0, ["Vanna client not initialized"]
    pairs = parse_sql_training_content(content)
    if not pairs:
        return 0, ["No question-SQL pairs parsed"]
    _ensure_sql_faiss_files_exist()
    trained_count = 0
    errors = []
    for i, pair in enumerate(pairs):
        try:
            vn.train(question=pair['question'], sql=pair['sql'])
            trained_count += 1
            logger.info(f"Trained SQL pair {i+1}/{len(pairs)}: {pair['question'][:50]}...")
        except Exception as e:
            error_msg = f"Failed to train pair {i+1}: {str(e)}"
            logger.error(error_msg)
            errors.append(error_msg)
    return trained_count, errors


async def _background_sync_sql_training():
    """
    Background task: sync all unsynced SQL training with Vanna.
    Fetches rows where is_sync=0, ensures sql files exist, trains, then sets is_sync=1.
    """
    try:
        from config import VANNA_ENABLED
        from services.semantic_layer import get_vanna_client
        if not VANNA_ENABLED:
            return
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, content, file_name 
                FROM sql_training 
                WHERE is_sync = 0
                ORDER BY id ASC
            """))
            rows = result.fetchall()
        if not rows:
            logger.info("Background SQL sync: no unsynced rows")
            return
        _ensure_sql_faiss_files_exist()
        vn = await get_vanna_client()
        if vn is None:
            logger.error("Background SQL sync: Vanna client not initialized")
            return
        total_trained = 0
        synced_ids = []
        for row in rows:
            doc_id, content, file_name = row[0], row[1], row[2]
            if not content or not content.strip():
                continue
            try:
                trained_count, _ = await _train_sql_content_with_vanna(content)
                total_trained += trained_count
                # Only mark as synced when at least one pair was trained (metadata/index updated)
                if trained_count > 0:
                    synced_ids.append(doc_id)
                else:
                    logger.warning(f"Background SQL sync: doc id {doc_id} produced 0 pairs; not marking synced")
            except Exception as e:
                logger.error(f"Background SQL sync: failed doc id {doc_id}: {e}")
        if synced_ids:
            placeholders = ",".join([str(i) for i in synced_ids])
            async with AsyncSessionLocal() as session:
                await session.execute(text(f"""
                    UPDATE sql_training SET is_sync = 1 WHERE id IN ({placeholders})
                """))
                await session.commit()
        logger.info(f"Background SQL sync: completed, synced {len(synced_ids)} document(s)")
    except Exception as e:
        logger.error(f"Background SQL sync failed: {e}")


@router.post("/sync/sql-training")
async def sync_sql_training():
    """
    Sync all unsynced SQL training with Vanna. Runs in background so frontend is not blocked.
    Returns immediately; training and is_sync=1 update happen asynchronously.
    """
    from config import VANNA_ENABLED
    
    if not VANNA_ENABLED:
        raise HTTPException(status_code=400, detail="Vanna is not enabled")
    
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id FROM sql_training WHERE is_sync = 0 ORDER BY id ASC
            """))
            rows = result.fetchall()
        
        if not rows:
            return {
                "success": True,
                "message": "No unsynced SQL training to sync",
                "trained_count": 0,
                "synced_ids": [],
                "background": False
            }
        
        asyncio.create_task(_background_sync_sql_training())
        
        return {
            "success": True,
            "message": "Sync started in background. Status will update when complete.",
            "trained_count": 0,
            "synced_ids": [r[0] for r in rows],
            "background": True
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error starting SQL sync: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to start SQL sync: {str(e)}")


@router.post("/sync/sql-training/{doc_id:int}")
async def sync_sql_training_single(doc_id: int):
    """
    Sync a single SQL training document with Vanna.
    Trains Vanna with that row's content and sets is_sync=1 for that row.
    """
    from config import VANNA_ENABLED
    
    if not VANNA_ENABLED:
        raise HTTPException(status_code=400, detail="Vanna is not enabled")
    
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, content, file_name FROM sql_training WHERE id = :doc_id
            """), {"doc_id": doc_id})
            row = result.fetchone()
        
        if not row:
            raise HTTPException(status_code=404, detail="SQL training document not found")
        
        content, file_name = row[1], row[2]
        if not content or not content.strip():
            return {
                "success": False,
                "message": "Document has no content to sync",
                "trained_count": 0
            }
        
        trained_count, errors = await _train_sql_content_with_vanna(content)
        
        # Only set is_sync=1 when at least one pair was trained (Vanna metadata/index updated)
        if trained_count > 0:
            async with AsyncSessionLocal() as session:
                await session.execute(text("UPDATE sql_training SET is_sync = 1 WHERE id = :doc_id"), {"doc_id": doc_id})
                await session.commit()
        
        from services.semantic_layer import get_vanna_client
        vn = await get_vanna_client()
        status = vn.get_training_status() if vn else {}
        
        return {
            "success": trained_count > 0,
            "message": f"Synced document (ID: {doc_id}, File: {file_name or 'text input'})" if trained_count > 0 else (errors[0] if errors else "No question-SQL pairs parsed from content"),
            "trained_count": trained_count,
            "errors": errors if errors else None,
            "training_status": status
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error syncing SQL training doc {doc_id}: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to sync document: {str(e)}")


async def _background_retrain_sql_training():
    """
    Background task: train Vanna with all remaining sql_training rows,
    then set is_sync=1 for all. Run after sql index was cleared (e.g. after delete).
    Recreates sql_metadata.json and sql_index.faiss only when rows exist; does not
    create them when the table is empty after deletion.
    """
    try:
        from config import VANNA_ENABLED, VANNA_FAISS_PATH
        from services.semantic_layer import get_vanna_client
        if not VANNA_ENABLED:
            return
        # Fetch rows first — do not create any sql files when no rows remain.
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, content FROM sql_training ORDER BY id ASC
            """))
            rows = result.fetchall()
        if not rows:
            logger.info("Background retrain SQL: no sql_training rows; not creating sql_metadata or sql_index")
            return
        # Rows present: ensure directory exists, then retrain (Vanna will create/write sql files).
        faiss_path = Path(VANNA_FAISS_PATH)
        faiss_path.mkdir(parents=True, exist_ok=True)
        vn = await get_vanna_client()
        if vn is None:
            logger.error("Background retrain SQL: Vanna client not initialized")
            return
        # Clear in-memory (and persisted) SQL training so retrain reflects only remaining rows.
        try:
            vn.remove_collection("sql")
        except Exception as e:
            logger.warning(f"Background retrain SQL: clear sql collection: {e}")
        for row in rows:
            doc_id, content = row[0], row[1]
            if not content or not content.strip():
                continue
            try:
                await _train_sql_content_with_vanna(content)
            except Exception as e:
                logger.error(f"Background retrain SQL: failed to train doc id {doc_id}: {e}")
        async with AsyncSessionLocal() as session:
            await session.execute(text("UPDATE sql_training SET is_sync = 1"))
            await session.commit()
        logger.info("Background retrain SQL: completed and set is_sync=1 for all")
    except Exception as e:
        logger.error(f"Background retrain sql_training failed: {e}")


@router.delete("/sql-training/{doc_id:int}")
async def delete_sql_training(doc_id: int):
    """
    Delete a SQL training document.
    Removes the row, sets all other rows is_sync=0, deletes sql_index.faiss and sql_metadata.json,
    then runs background retrain with remaining rows and sets is_sync=1.
    Response returns immediately; retrain runs async in background.
    """
    from config import VANNA_FAISS_PATH
    
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("SELECT id FROM sql_training WHERE id = :doc_id"), {"doc_id": doc_id})
            if result.fetchone() is None:
                raise HTTPException(status_code=404, detail="SQL training document not found")
            await session.execute(text("DELETE FROM sql_training WHERE id = :doc_id"), {"doc_id": doc_id})
            await session.execute(text("UPDATE sql_training SET is_sync = 0"))
            await session.commit()
        
        faiss_path = Path(VANNA_FAISS_PATH)
        sql_index = faiss_path / "sql_index.faiss"
        sql_metadata = faiss_path / "sql_metadata.json"
        for f in (sql_index, sql_metadata):
            if f.exists():
                try:
                    f.unlink()
                    logger.info(f"Deleted {f.name} for SQL training retrain")
                except Exception as e:
                    logger.warning(f"Could not delete {f}: {e}")
        
        asyncio.create_task(_background_retrain_sql_training())
        
        return {
            "success": True,
            "message": "Document deleted. Vanna is retraining with remaining SQL documents in the background.",
            "deleted_id": doc_id
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting SQL training: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to delete document: {str(e)}")


async def _background_retrain_business_context_docs():
    """
    Background task: train Vanna with all remaining business_context rows,
    then set is_sync=1 for all. Run after doc index was cleared (e.g. after delete).
    Clears in-memory documentation collection first so metadata/index match remaining rows only.
    When there are no rows, returns without calling remove_collection so empty
    documentation_index.faiss / documentation_metadata.json are not created.
    """
    try:
        # Fetch remaining rows first; if none, exit without touching Vanna (avoids creating empty doc index files).
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("""
                SELECT id, content FROM business_context ORDER BY id ASC
            """))
            rows = result.fetchall()
        if not rows:
            logger.info("Background retrain: no business context rows to train")
            return
        from config import VANNA_ENABLED
        from services.semantic_layer import get_vanna_client
        if not VANNA_ENABLED:
            return
        vn = await get_vanna_client()
        if vn is None:
            logger.error("Background retrain: Vanna client not initialized")
            return
        # Clear in-memory (and persisted) documentation so retrain reflects only remaining rows.
        try:
            vn.remove_collection("documentation")
        except Exception as e:
            logger.warning(f"Background retrain: clear documentation collection: {e}")
        for row in rows:
            doc_id, content = row[0], row[1]
            if not content or not content.strip():
                continue
            try:
                await _train_content_with_vanna(content)
            except Exception as e:
                logger.error(f"Background retrain: failed to train doc id {doc_id}: {e}")
        async with AsyncSessionLocal() as session:
            await session.execute(text("UPDATE business_context SET is_sync = 1"))
            await session.commit()
        logger.info("Background retrain: completed and set is_sync=1 for all")
    except Exception as e:
        logger.error(f"Background retrain business context failed: {e}")


@router.delete("/business-context/{doc_id:int}")
async def delete_business_context(doc_id: int):
    """
    Delete a business context document.
    Removes the row, sets all other rows is_sync=0, deletes doc_index.faiss and doc_metadata.json,
    then runs background retrain with remaining rows and sets is_sync=1.
    Response returns immediately; retrain runs async in background.
    """
    from config import VANNA_FAISS_PATH
    
    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(text("SELECT id FROM business_context WHERE id = :doc_id"), {"doc_id": doc_id})
            if result.fetchone() is None:
                raise HTTPException(status_code=404, detail="Business context document not found")
            await session.execute(text("DELETE FROM business_context WHERE id = :doc_id"), {"doc_id": doc_id})
            await session.execute(text("UPDATE business_context SET is_sync = 0"))
            await session.commit()
            # Count remaining rows; if none, skip background retrain so empty doc index files are not created
            count_result = await session.execute(text("SELECT COUNT(*) FROM business_context"))
            remaining = count_result.scalar() or 0
        if remaining == 0:
            # Delete doc FAISS index files when last row removed; do not run background retrain
            faiss_path = Path(VANNA_FAISS_PATH)
            for name in ("doc_index.faiss", "doc_metadata.json", "documentation_index.faiss", "documentation_metadata.json"):
                f = faiss_path / name
                if f.exists():
                    try:
                        f.unlink()
                        logger.info(f"Deleted {f.name} (last business context row removed)")
                    except Exception as e:
                        logger.warning(f"Could not delete {f}: {e}")
        else:
            # Delete doc FAISS index files so Vanna rebuilds from remaining rows
            faiss_path = Path(VANNA_FAISS_PATH)
            doc_index = faiss_path / "doc_index.faiss"
            doc_metadata = faiss_path / "doc_metadata.json"
            for f in (doc_index, doc_metadata):
                if f.exists():
                    try:
                        f.unlink()
                        logger.info(f"Deleted {f.name} for business context retrain")
                    except Exception as e:
                        logger.warning(f"Could not delete {f}: {e}")
            # Run retrain in background (do not await)
            asyncio.create_task(_background_retrain_business_context_docs())
        
        return {
            "success": True,
            "message": "Document deleted. Vanna is retraining with remaining documents in the background." if remaining > 0 else "Document deleted. No business context documents remain; doc index files removed.",
            "deleted_id": doc_id
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting business context: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to delete document: {str(e)}")


@router.get("/status")
async def get_training_status():
    """Get current Vanna training status."""
    from config import VANNA_ENABLED
    
    if not VANNA_ENABLED:
        return {
            "success": True,
            "vanna_enabled": False,
            "message": "Vanna is not enabled"
        }
    
    try:
        from services.semantic_layer import get_vanna_client
        
        vn = await get_vanna_client()
        if vn is None:
            return {
                "success": False,
                "vanna_enabled": True,
                "message": "Vanna client not initialized"
            }
        
        status = vn.get_training_status()
        
        return {
            "success": True,
            "vanna_enabled": True,
            "training_status": status
        }
        
    except Exception as e:
        logger.error(f"Error getting training status: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to get training status: {str(e)}")
